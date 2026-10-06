"""
One-time local import: converts the CA's unit-master Excel files (Sun Atmosphere,
Sun Gravitas) into data/unit_master_<project_code>.json, in the same shape the
app already auto-loads at startup (mirrors data/synced_sun_footprint.json).

These files only contain unit inventory (block/flat/size) - no owner or payment
data. Ownership & payment history is derived separately at render time from the
already-synced Tally voucher records, matched by unit_no.

Usage: py import_unit_master.py
"""
import os
import json
import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
os.makedirs(DATA_DIR, exist_ok=True)


def safe_float(val):
    try:
        if val is None or str(val).strip() == "":
            return 0.0
        return float(val)
    except (ValueError, TypeError):
        return 0.0


def norm_unit_no(block, flat):
    b = str(block).strip().upper() if block not in (None, "") else ""
    f = str(flat).strip()
    if isinstance(flat, float) and flat.is_integer():
        f = str(int(flat))
    return f"{b}-{f}" if b else f


def import_atmosphere(path, project_code, project_name):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))
    units = []
    for r in rows[1:]:
        if not r or r[1] is None or r[2] is None:
            continue
        block, flat, usage, size_sqft, carpet_sqft, terrace_sqft = r[1], r[2], r[3], r[4], r[5], r[6]
        units.append({
            "unit_no": norm_unit_no(block, flat),
            "block": str(block).strip().upper(),
            "flat_label": str(flat).strip(),
            "floor": "",
            "usage_type": (usage or "").strip(),
            "office_no": "",
            "size_sqft": safe_float(size_sqft),
            "carpet_sqft": safe_float(carpet_sqft),
            "terrace_sqft": safe_float(terrace_sqft),
        })
    return {"project": project_name, "company_code": project_code, "count": len(units), "units": units}


def import_gravitas(path, project_code, project_name):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))
    units = []
    for r in rows[3:]:
        if not r or r[2] is None or r[3] is None:
            continue
        if str(r[0]).strip().lower() == "total":
            continue
        floor, block, unit_no, office_no, size_sqft = r[1], r[2], r[3], r[4], r[5]
        units.append({
            "unit_no": norm_unit_no(block, unit_no),
            "block": str(block).strip().upper(),
            "flat_label": str(unit_no).strip(),
            "floor": str(floor).strip() if floor else "",
            "usage_type": "Commercial",
            "office_no": str(office_no).strip() if office_no else "",
            "size_sqft": safe_float(size_sqft),
            "carpet_sqft": 0.0,
            "terrace_sqft": 0.0,
        })
    return {"project": project_name, "company_code": project_code, "count": len(units), "units": units}


if __name__ == "__main__":
    atmosphere = import_atmosphere(r"C:\naman_ca\SUN ATmospher- Project details.xlsx", "010000", "Sun Atmosphere")
    gravitas = import_gravitas(r"C:\naman_ca\SUN Gravitas- Project Details.xlsx", "010009", "Sun Gravitas")

    with open(os.path.join(DATA_DIR, "unit_master_010000.json"), "w", encoding="utf-8") as f:
        json.dump(atmosphere, f, indent=2)
    with open(os.path.join(DATA_DIR, "unit_master_010009.json"), "w", encoding="utf-8") as f:
        json.dump(gravitas, f, indent=2)

    print(f"Sun Atmosphere: {atmosphere['count']} units -> data/unit_master_010000.json")
    print(f"Sun Gravitas:   {gravitas['count']} units -> data/unit_master_010009.json")
