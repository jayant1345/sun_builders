import openpyxl
import xml.etree.ElementTree as ET
import os

class TallyFileReader:
    """
    Parses offline exports from Tally (Daybook / Receipts Register in XML or Excel format).
    """
    @staticmethod
    def read_daybook_xml(xml_file_path: str) -> list:
        vouchers = []
        if not os.path.exists(xml_file_path):
            return vouchers

        try:
            tree = ET.parse(xml_file_path)
            root = tree.getroot()
            for vch in root.findall(".//VOUCHER"):
                v_type = vch.findtext("VOUCHERTYPENAME", "")
                v_date = vch.findtext("DATE", "")
                v_num = vch.findtext("VOUCHERNUMBER", "")
                narration = vch.findtext("NARRATION", "")

                entries = []
                for le in vch.findall(".//ALLLEDGERENTRIES.LIST"):
                    name = le.findtext("LEDGERNAME", "")
                    amt = float(le.findtext("AMOUNT", "0") or 0)
                    entries.append({"ledger_name": name, "amount": amt})

                vouchers.append({
                    "voucher_type": v_type,
                    "date": v_date,
                    "voucher_number": v_num,
                    "narration": narration,
                    "entries": entries
                })
        except Exception as e:
            print(f"[TallyFileReader] Error parsing XML file: {e}")
        return vouchers

    @staticmethod
    def read_receipts_excel(excel_file_path: str) -> list:
        records = []
        if not os.path.exists(excel_file_path):
            return records

        try:
            wb = openpyxl.load_workbook(excel_file_path, data_only=True)
            ws = wb.active
            for r in range(2, ws.max_row + 1):
                date_val = str(ws.cell(r, 1).value or "")
                ledger_name = str(ws.cell(r, 2).value or "")
                amt = ws.cell(r, 3).value
                if isinstance(amt, (int, float)) and amt > 0:
                    records.append({
                        "date": date_val,
                        "ledger_name": ledger_name,
                        "amount": amt
                    })
        except Exception as e:
            print(f"[TallyFileReader] Error parsing Excel dump: {e}")
        return records
