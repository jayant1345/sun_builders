import os
import re
import shutil
import calendar
from datetime import datetime
import openpyxl
from core.gst_rules import is_creditor_party, is_after_bu_member

class ExcelGenerator:
    """
    Generates CA-standard multi-sheet GSTR-1 Excel files with live formulas.
    """
    def __init__(self, template_path: str):
        self.template_path = template_path

    @staticmethod
    def parse_month_year(target_month_str: str):
        """
        Parses a month string like 'SEP-26', 'SEP 2026', 'SEPTEMBER-2026'
        into (month_abbr, month_full, short_year, full_year, month_num, last_day).
        """
        months_map = {
            "JAN": (1, "JANUARY"), "FEB": (2, "FEBRUARY"), "MAR": (3, "MARCH"),
            "APR": (4, "APRIL"), "MAY": (5, "MAY"), "JUN": (6, "JUNE"),
            "JUL": (7, "JULY"), "AUG": (8, "AUGUST"), "SEP": (9, "SEPTEMBER"),
            "OCT": (10, "OCTOBER"), "NOV": (11, "NOVEMBER"), "DEC": (12, "DECEMBER")
        }
        s = (target_month_str or "SEP-26").strip().upper()
        
        # Find month
        m_num = 9
        m_abbr = "SEP"
        m_full = "SEPTEMBER"
        for k, (num, full) in months_map.items():
            if k in s or full in s:
                m_num = num
                m_abbr = k
                m_full = full
                break

        # Find year
        y_match = re.search(r'20(\d{2})', s)
        if y_match:
            short_y = y_match.group(1)
            full_y = "20" + short_y
        else:
            y2 = re.search(r'[\'\-_/\s](\d{2})\b', s)
            if y2:
                short_y = y2.group(1)
                full_y = "20" + short_y
            else:
                short_y = "26"
                full_y = "2026"

        last_day = calendar.monthrange(int(full_y), m_num)[1]
        return m_abbr, m_full, short_y, full_y, m_num, last_day

    def generate_monthly_workbook(self, target_month_str: str, project_data: dict, output_path: str) -> str:
        """
        Creates a new workbook from template, updates month headers, injects unit-wise
        collections for the target tax period, and saves to output_path with intact formulas.
        """
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        # Copy template directly to preserve exact styles, merged cells, fonts, and formulas
        shutil.copy2(self.template_path, output_path)

        wb = openpyxl.load_workbook(output_path)

        m_abbr, m_full, short_y, full_y, m_num, last_day = self.parse_month_year(target_month_str)
        target_display = f"{m_abbr}-{short_y}"
        target_display_tick = f"{m_abbr}'{short_y}"
        target_display_space = f"{m_abbr} - {short_y}"
        last_date_obj = datetime(int(full_y), m_num, last_day)

        # 1. Update Month headers, labels & dates across all sheets
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    val = cell.value
                    if isinstance(val, datetime):
                        # Update month-end date for JV entries
                        if val.year == 2026 or val.year == int(full_y):
                            cell.value = last_date_obj
                    elif isinstance(val, str) and val.strip():
                        # Replace August 2026 patterns with target month
                        new_val = val
                        new_val = re.sub(r"AUG\s*'\s*26\b", target_display_tick, new_val, flags=re.IGNORECASE)
                        new_val = re.sub(r'AUG[\s\-_]+26\b', target_display, new_val, flags=re.IGNORECASE)
                        new_val = re.sub(r'AUG26\b', target_display, new_val, flags=re.IGNORECASE)
                        new_val = re.sub(r'AUGUST[\s\-_]+2026\b', f"{m_full} {full_y}", new_val, flags=re.IGNORECASE)
                        new_val = re.sub(r'AUGUST[\s\-_]+26\b', f"{m_full}-{short_y}", new_val, flags=re.IGNORECASE)
                        if new_val != val:
                            cell.value = new_val

        # 2. Populate or Refresh Member Collection Sheets
        # Mapping known projects to their sheets
        # 010010: "SUN REALTY FOOTPRINT- MEMBERS"
        # 010011: "SUN PARK WEST- MEMBER COLL."
        # 010000: "DATA"
        # 010009: "SUN REALTY-MEM COLL"
        # 010002: "SUN SILVER SPRING-MEMBER COLL"
        # 010015: "LEKHAMBHA- Member coll"
        
        self._populate_footprint_sheet(wb, project_data.get("010010") or project_data.get("FOOTPRINT") or [])
        self._populate_parkwest_sheet(wb, project_data.get("010011") or project_data.get("PARK WEST") or [])

        wb.save(output_path)
        print(f"[ExcelGenerator] Generated finalized GST workbook: {output_path}")
        return output_path

    def _populate_footprint_sheet(self, wb, vouchers):
        """
        Populates SUN REALTY FOOTPRINT- MEMBERS with actual target month vouchers.
        Follows CA directives:
        1. Creditors/vendors strictly excluded.
        2. Flats purchased after BU permission are 0% GST (Completely Exempt).
        3. Tax on Before-BU advances calculated considering the receipt as inclusive of tax base.
        """
        sheet_name = "SUN REALTY FOOTPRINT- MEMBERS"
        if sheet_name not in wb.sheetnames:
            return
        ws = wb[sheet_name]

        # 1. Clear residential rows (9 to 217)
        for r in range(9, 218):
            ws.cell(r, 1).value = None
            ws.cell(r, 2).value = None
            ws.cell(r, 3).value = None
            ws.cell(r, 4).value = None
            ws.cell(r, 5).value = f'=+D{r}'
            ws.cell(r, 6).value = None
            ws.cell(r, 7).value = f'=+F{r}'
            ws.cell(r, 8).value = None
            ws.cell(r, 9).value = None
            ws.cell(r, 10).value = None
            ws.cell(r, 11).value = None
            ws.cell(r, 12).value = None
            ws.cell(r, 13).value = f'=C{r}-D{r}-F{r}-H{r}'
            ws.cell(r, 14).value = f'=ROUND(M{r}*0.5/101, 2)'

        # 2. Clear shop rows (223 to 229)
        for r in range(223, 230):
            ws.cell(r, 1).value = None
            ws.cell(r, 2).value = None
            ws.cell(r, 3).value = None
            ws.cell(r, 4).value = None
            ws.cell(r, 5).value = f'=+D{r}'
            ws.cell(r, 6).value = None
            ws.cell(r, 7).value = f'=+F{r}'
            ws.cell(r, 8).value = None
            ws.cell(r, 9).value = None
            ws.cell(r, 10).value = None
            ws.cell(r, 11).value = None
            ws.cell(r, 12).value = None
            ws.cell(r, 13).value = f'=C{r}-D{r}-F{r}'
            ws.cell(r, 14).value = f'=ROUND(M{r}*2.5/105, 2)'

        if not vouchers:
            return

        # 3. Group vouchers by unit (strictly excluding creditors)
        res_units = {}
        shop_units = {}
        total_after_bu_exempt = 0.0

        for v in vouchers:
            party = str(v.get("party_name") or v.get("name") or v.get("member_name") or "").strip()
            orig = str(v.get("party_original") or v.get("raw_name") or "").strip()
            narr = str(v.get("narration") or "").strip()

            # Strictly exclude creditors, vendors, suppliers, contractors and pass-through bank entries
            if is_creditor_party(party) or is_creditor_party(orig) or is_creditor_party(narr):
                continue
            if not v.get("is_income", True):
                continue

            unit = str(v.get("unit_no") or v.get("unit") or v.get("flat_no") or "").strip()
            amt = float(v.get("cr_amount") or v.get("amount") or 0)
            reg_fee = float(v.get("reg_fee") or 0)
            stamp_duty = float(v.get("stamp_duty") or 0)

            unit_is_after_bu = is_after_bu_member(party, orig, narr)
            if unit_is_after_bu:
                total_after_bu_exempt += max(0.0, amt - reg_fee - stamp_duty)

            # Check if residential or commercial
            if re.match(r'^[A-Za-z]\s*[\-_/]', unit):
                if unit not in res_units:
                    res_units[unit] = {"unit": unit, "party": party, "amount": 0.0, "reg_fee": 0.0, "stamp_duty": 0.0, "is_after_bu": unit_is_after_bu}
                res_units[unit]["amount"] += amt
                res_units[unit]["reg_fee"] += reg_fee
                res_units[unit]["stamp_duty"] += stamp_duty
                if party and not res_units[unit]["party"]:
                    res_units[unit]["party"] = party
                if unit_is_after_bu:
                    res_units[unit]["is_after_bu"] = True
            else:
                # Genuine shop/commercial units only
                if not unit and not any(k in party.lower() for k in ["shop", "sh-", "sh "]):
                    continue
                u_key = unit or f"Shop-{party[:15]}"
                if u_key not in shop_units:
                    shop_units[u_key] = {"unit": unit or "", "party": party, "amount": 0.0, "reg_fee": 0.0, "stamp_duty": 0.0, "is_after_bu": unit_is_after_bu}
                shop_units[u_key]["amount"] += amt
                shop_units[u_key]["reg_fee"] += reg_fee
                shop_units[u_key]["stamp_duty"] += stamp_duty
                if party and not shop_units[u_key]["party"]:
                    shop_units[u_key]["party"] = party
                if unit_is_after_bu:
                    shop_units[u_key]["is_after_bu"] = True

        # 4. Populate residential units into rows 9..217
        r_idx = 9
        for u, data in sorted(res_units.items()):
            if r_idx >= 218:
                break
            ws.cell(r_idx, 1).value = data["unit"]
            ws.cell(r_idx, 2).value = data["party"]
            ws.cell(r_idx, 3).value = data["amount"]
            if data["reg_fee"] > 0:
                ws.cell(r_idx, 4).value = data["reg_fee"]
            if data["stamp_duty"] > 0:
                ws.cell(r_idx, 6).value = data["stamp_duty"]

            # Tax Calculation:
            if data.get("is_after_bu"):
                # Flats purchased after BU permission are 0% GST (Completely Exempt)
                ws.cell(r_idx, 13).value = 0
                ws.cell(r_idx, 14).value = 0
            else:
                # Before-BU advances: Tax is calculated inclusive of tax base (1% rate: CGST 0.5%)
                ws.cell(r_idx, 13).value = f'=C{r_idx}-D{r_idx}-F{r_idx}-H{r_idx}'
                ws.cell(r_idx, 14).value = f'=ROUND(M{r_idx}*0.5/101, 2)'
            r_idx += 1

        # 5. Populate shop units into rows 223..229
        s_idx = 223
        for u, data in sorted(shop_units.items()):
            if s_idx >= 230:
                break
            ws.cell(s_idx, 1).value = data["unit"]
            ws.cell(s_idx, 2).value = data["party"]
            ws.cell(s_idx, 3).value = data["amount"]
            if data["reg_fee"] > 0:
                ws.cell(s_idx, 4).value = data["reg_fee"]
            if data["stamp_duty"] > 0:
                ws.cell(s_idx, 6).value = data["stamp_duty"]

            if data.get("is_after_bu"):
                ws.cell(s_idx, 13).value = 0
                ws.cell(s_idx, 14).value = 0
            else:
                # Commercial @ 5% inclusive (CGST 2.5%)
                ws.cell(s_idx, 13).value = f'=C{s_idx}-D{s_idx}-F{s_idx}'
                ws.cell(s_idx, 14).value = f'=ROUND(M{s_idx}*2.5/105, 2)'
            s_idx += 1

        # 6. Update After-BU Exempt Sale in FOOTPRINT- CALCULATION Row 17
        if "FOOTPRINT- CALCULATION" in wb.sheetnames:
            ws_calc = wb["FOOTPRINT- CALCULATION"]
            if total_after_bu_exempt > 0:
                ws_calc.cell(17, 2).value = round(total_after_bu_exempt, 2)
            # Ensure Row 6 formula computes tax inclusive of tax base
            ws_calc.cell(6, 4).value = "=ROUND(B6*0.5/101, 2)"
            ws_calc.cell(6, 5).value = "=+D6"

    def _populate_parkwest_sheet(self, wb, vouchers):
        """
        Populates SUN PARK WEST- MEMBER COLL. with target month vouchers.
        """
        sheet_name = "SUN PARK WEST- MEMBER COLL."
        if sheet_name not in wb.sheetnames:
            return
        ws = wb[sheet_name]

        # Clear residential rows 9..73
        for r in range(9, 74):
            ws.cell(r, 1).value = None
            ws.cell(r, 2).value = None
            ws.cell(r, 3).value = None
            ws.cell(r, 4).value = None
            ws.cell(r, 5).value = f'=+D{r}'
            ws.cell(r, 6).value = None
            ws.cell(r, 7).value = f'=+F{r}'
            ws.cell(r, 8).value = None
            ws.cell(r, 9).value = None

        if not vouchers:
            return

        # Group and populate
        units_map = {}
        for v in vouchers:
            party = str(v.get("party_name") or v.get("name") or v.get("member_name") or "").strip()
            orig = str(v.get("party_original") or v.get("raw_name") or "").strip()
            if is_creditor_party(party) or is_creditor_party(orig):
                continue
            if not v.get("is_income", True):
                continue
            unit = str(v.get("unit_no") or v.get("unit") or v.get("flat_no") or "").strip()
            amt = float(v.get("cr_amount") or v.get("amount") or 0)
            if unit not in units_map:
                units_map[unit] = {"unit": unit, "party": party, "amount": 0.0}
            units_map[unit]["amount"] += amt

        r_idx = 9
        for u, data in sorted(units_map.items()):
            if r_idx >= 74:
                break
            ws.cell(r_idx, 1).value = data["unit"]
            ws.cell(r_idx, 2).value = data["party"]
            ws.cell(r_idx, 3).value = data["amount"]
            r_idx += 1
